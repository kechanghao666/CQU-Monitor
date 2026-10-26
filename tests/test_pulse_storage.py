import csv
import json
import sqlite3
import threading

import pytest
from openpyxl import load_workbook

from cqu_monitor.storage import QueuedRecorder, SessionStore, StorageError


START = "2026-10-16T01:00:00+08:00"


def observation(sid, elapsed=0, group="left", sequence=1):
    return dict(
        session_id=sid,
        elapsed=elapsed,
        wall_time=START,
        group=group,
        sequence=sequence,
        started=elapsed + 10,
        finished=elapsed + 10.1,
        raw_parameters=[600, 0, 0, 15],
        raw_registers=[1, 380, 0, 0, 2, 0],
        width=6.0,
        outputs=15,
        channel=1,
        remaining=3.8,
        cycle=2,
        interval_remaining=0,
        cycle_remaining=0,
        status="active",
        error=None,
        extra_audit={"firmware": "unknown"},
    )


def sample(sid, elapsed=0):
    return dict(
        session_id=sid,
        elapsed=elapsed,
        wall_time=START,
        cell=1,
        role="voltage_left",
        valid=True,
        value=1.2,
        synchronization_source="controller",
        sync_evidence={"before": {"sequence": 1}, "after": {"sequence": 2}},
    )


@pytest.fixture
def store(tmp_path):
    return SessionStore(tmp_path / "history.sqlite3")


def create(store):
    return store.create_session("hardware", {"synchronization": "controller"}, START)


def test_control_roundtrip_filters_pagination_and_schema(store):
    sid, other = create(store), create(store)
    records = [observation(sid, i, "left" if i % 2 == 0 else "right", i + 1) for i in range(5)]
    records[0]["wall_time"] = "2026-10-15T17:00:00+00:00"
    store.write_batch(controls=[*reversed(records), observation(other)])
    assert store.query_controls(sid) == records
    assert store.query_controls(sid, start=1, end=4, group="right", limit=1, offset=1) == [
        records[3]
    ]
    assert store.query_controls(sid, start=START, end=START) == records
    assert store.query_controls("missing") == []
    for bad in ("LEFT", "left' OR 1=1 --", "current", 0):
        with pytest.raises(StorageError, match="left.*right"):
            store.query_controls(sid, group=bad)
    for kwargs in ({"limit": -1}, {"offset": -1}, {"limit": 1.2}):
        with pytest.raises(StorageError):
            store.query_controls(sid, **kwargs)
    with sqlite3.connect(store.db_path) as conn:
        assert conn.execute("PRAGMA user_version").fetchone()[0] == 3
        indexes = conn.execute("PRAGMA index_list(control_observations)").fetchall()
        assert any(
            [r[2] for r in conn.execute(f"PRAGMA index_info({i[1]})")]
            == ["session_id", "elapsed", "id"]
            for i in indexes
        )


def test_v2_migration_preserves_payload_text_settings_snapshot_and_ids(tmp_path):
    path = tmp_path / "v2.sqlite3"
    raw_sample = '{ "session_id": "old", "elapsed": 2, "custom": [1, 2] }'
    settings, snapshot = '{ "mode": "old", "custom": 42 }', '{ "cells": [] }'
    with sqlite3.connect(path) as conn:
        conn.executescript("""
            CREATE TABLE sessions(id TEXT PRIMARY KEY,mode TEXT NOT NULL,
              started_at TEXT NOT NULL,ended_at TEXT,status TEXT NOT NULL,
              settings TEXT NOT NULL,snapshot TEXT);
            CREATE TABLE samples(id INTEGER PRIMARY KEY AUTOINCREMENT,
              session_id TEXT NOT NULL REFERENCES sessions(id),elapsed REAL,
              wall_time TEXT NOT NULL,cell INTEGER,role TEXT NOT NULL,
              valid INTEGER NOT NULL,payload TEXT NOT NULL);
            CREATE TABLE summaries(id INTEGER PRIMARY KEY AUTOINCREMENT,
              session_id TEXT NOT NULL REFERENCES sessions(id),elapsed REAL NOT NULL,
              wall_time TEXT NOT NULL,payload TEXT NOT NULL);
            CREATE TABLE events(id INTEGER PRIMARY KEY AUTOINCREMENT,
              session_id TEXT NOT NULL REFERENCES sessions(id),
              wall_time TEXT NOT NULL,payload TEXT NOT NULL);
            PRAGMA user_version=2;
        """)
        conn.execute(
            "INSERT INTO sessions VALUES(?,?,?,?,?,?,?)",
            ("old", "hardware", START, START, "completed", settings, snapshot),
        )
        conn.execute(
            "INSERT INTO samples VALUES(?,?,?,?,?,?,?,?)",
            (19, "old", 2, START, 1, "voltage_left", 1, raw_sample),
        )
        conn.execute("INSERT INTO summaries VALUES(?,?,?,?,?)", (8, "old", 2, START, '{ "x": 4 }'))
        conn.execute(
            "INSERT INTO events VALUES(?,?,?,?)", (12, "old", START, '{ "level": "info" }')
        )
        before = {
            table: conn.execute(f"SELECT * FROM {table}").fetchall()
            for table in ("sessions", "samples", "summaries", "events")
        }
        roots = dict(conn.execute("SELECT name,rootpage FROM sqlite_master WHERE type='table'"))
    migrated = SessionStore(path)
    assert migrated.query_controls("old") == []
    with sqlite3.connect(path) as conn:
        for table, rows in before.items():
            assert conn.execute(f"SELECT * FROM {table}").fetchall() == rows
        after_roots = dict(
            conn.execute("SELECT name,rootpage FROM sqlite_master WHERE type='table'")
        )
        assert all(after_roots[name] == page for name, page in roots.items())
    migrated.write_batch(controls=[observation("old")])
    assert len(SessionStore(path).query_controls("old")) == 1


def test_control_failure_rolls_back_samples_summaries_events_snapshot_and_controls(store):
    sid = create(store)
    before = {"session_id": sid, "phase": "before"}
    store.write_batch(snapshot=before)
    summary = dict(session_id=sid, elapsed=0, wall_time=START)
    event = dict(session_id=sid, wall_time=START, level="info", message="audit")
    with pytest.raises(StorageError):
        store.write_batch(
            samples=[sample(sid)],
            summaries=[summary],
            events=[event],
            controls=[observation(sid), observation("missing")],
            snapshot={"session_id": sid, "phase": "after"},
        )
    assert store.query_samples(sid) == store.query_summaries(sid) == store.query_controls(sid) == []
    assert store.get_session(sid)["snapshot"] == before
    with sqlite3.connect(store.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM events").fetchone()[0] == 0
    with pytest.raises(StorageError):
        store.write_batch(controls=[observation(sid)], snapshot={"session_id": "missing"})
    assert store.query_controls(sid) == []


def test_recorder_drains_controls_freezes_payload_and_validates_session(store):
    sid = create(store)
    recorder = QueuedRecorder(store, sid, batch_size=50, flush_interval=30)
    recorder.start()
    row = observation(sid)
    try:
        with pytest.raises(StorageError, match="编号"):
            recorder.submit("control", observation("missing"))
        recorder.submit("control", row)
        row["raw_registers"][0] = 99
        for i in range(1, 14):
            recorder.submit("control", observation(sid, i, sequence=i + 1))
    finally:
        recorder.close({"session_id": sid}, START)
    rows = store.query_controls(sid)
    assert len(rows) == 14 and rows[0]["raw_registers"][0] == 1
    assert recorder.pending_count == 0 and not recorder.alive
    assert store.get_session(sid)["status"] == "completed"


def test_failed_controls_preserve_inflight_and_remaining_queue(store, monkeypatch):
    sid = create(store)
    entered, release = threading.Event(), threading.Event()

    def fail(**kwargs):
        assert len(kwargs["controls"]) == 1
        entered.set()
        assert release.wait(2)
        raise OSError("injected control disk full")

    monkeypatch.setattr(store, "write_batch", fail)
    recorder = QueuedRecorder(store, sid, batch_size=1)
    recorder.start()
    try:
        recorder.submit("control", observation(sid, sequence=1))
        assert entered.wait(2)
        for i in range(2, 9):
            recorder.submit("control", observation(sid, i, sequence=i))
    finally:
        release.set()
        with pytest.raises(StorageError, match="control disk full"):
            recorder.close({"session_id": sid}, START)
    records = [
        json.loads(line) for line in recorder.recovery_path.read_text(encoding="utf-8").splitlines()
    ]
    assert [r["payload"]["sequence"] for r in records if r["kind"] == "control"] == list(
        range(1, 9)
    )
    assert recorder.pending_count == 8 and not recorder.alive
    assert store.query_controls(sid) == []
    assert store.get_session(sid)["status"] == "running"


def exported_tables(path):
    if path.suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return {name: list(wb[name].values) for name in wb.sheetnames}
        finally:
            wb.close()
    rows = list(csv.reader(path.open(encoding="utf-8-sig", newline=""), delimiter="\t"))
    tables, current = {}, None
    for row in rows:
        if len(row) == 1 and row[0].startswith("["):
            current = row[0][1:-1]
            tables[current] = []
        elif row and current:
            tables[current].append(tuple(row))
    return tables


@pytest.mark.parametrize("suffix", [".txt", ".xlsx"])
def test_exports_keep_named_raw_fields_full_payload_and_sync_evidence(store, tmp_path, suffix):
    sid = create(store)
    control = observation(sid)
    control["status"], control["error"] = "=unsafe()", "=error()"
    store.write_batch(samples=[sample(sid)], controls=[control])
    path = tmp_path / f"audit{suffix}"
    store.export_session(sid, path)
    tables = exported_tables(path)
    headers, row = tables["控制器状态"]
    exported = dict(zip(headers, row))
    assert set(control) - {"extra_audit"} <= set(headers)
    assert json.loads(exported["raw_parameters"]) == control["raw_parameters"]
    assert json.loads(exported["raw_registers"]) == control["raw_registers"]
    assert json.loads(exported["payload"]) == control
    assert exported["status"] == ("'=unsafe()" if suffix == ".xlsx" else "=unsafe()")
    headers, row = tables["采样数据"]
    exported = dict(zip(headers, row))
    assert exported["synchronization_source"] == "controller"
    assert json.loads(exported["sync_evidence"]) == sample(sid)["sync_evidence"]


@pytest.mark.parametrize(
    "cell,expected", [(None, [1, 2]), (1, [1]), (15, [1]), (16, [2]), (30, [2]), (0, [1, 2])]
)
def test_export_control_time_and_group_filters_keep_errors_and_current_context(
    store, tmp_path, cell, expected
):
    sid = create(store)
    rows = [
        observation(sid, 2, "left", 1),
        observation(sid, 3, "right", 2),
        observation(sid, 4, "left", 3),
    ]
    rows[0].update(status="error", error="disconnected")
    store.write_batch(controls=rows)
    path = tmp_path / "filtered.txt"
    store.export_session(sid, path, start=2, end=3, cell=cell, valid_only=True)
    headers, *values = exported_tables(path)["控制器状态"]
    assert [int(row[headers.index("sequence")]) for row in values] == expected
    # 电流独立采集；筛选电流仍保留同期左右控制器上下文以供追溯。
    datepath = tmp_path / "date.txt"
    store.export_session(sid, datepath, start="2026-10-15T17:00:00+00:00", end=START)
    assert len(exported_tables(datepath)["控制器状态"]) == 4


def test_export_empty_controls_old_sample_blanks_and_controller_excel_split(
    store, tmp_path, monkeypatch
):
    import cqu_monitor.storage as storage

    sid = create(store)
    row = sample(sid)
    row.pop("synchronization_source")
    row.pop("sync_evidence")
    store.write_batch(samples=[row])
    empty = tmp_path / "old.xlsx"
    store.export_session(sid, empty)
    tables = exported_tables(empty)
    assert len(tables["控制器状态"]) == 1
    headers, oldrow = tables["采样数据"]
    assert oldrow[headers.index("sync_evidence")] is None
    assert oldrow[headers.index("synchronization_source")] is None
    monkeypatch.setattr(storage, "EXCEL_MAX_ROWS", 4)
    store.write_batch(controls=[observation(sid, i, sequence=i) for i in range(7)])
    path = tmp_path / "split.xlsx"
    store.export_session(sid, path)
    tables = exported_tables(path)
    controls = [rows for name, rows in tables.items() if name.startswith("控制器状态")]
    assert len(controls) == 3
    assert sum(len(rows) - 1 for rows in controls) == 7


def test_recorder_without_controls_keeps_legacy_batch_signature(store, monkeypatch):
    sid = create(store)
    original = store.write_batch

    def legacy_write(samples=(), summaries=(), events=(), snapshot=None):
        original(samples=samples, summaries=summaries, events=events, snapshot=snapshot)

    monkeypatch.setattr(store, "write_batch", legacy_write)
    recorder = QueuedRecorder(store, sid)
    recorder.start()
    recorder.submit("sample", sample(sid))
    recorder.close({"session_id": sid}, START)
    assert len(store.query_samples(sid)) == 1
    assert recorder.pending_count == 0


def test_control_invalid_json_rejected_before_accepting_queue_item(store):
    sid = create(store)
    recorder = QueuedRecorder(store, sid, flush_interval=30)
    recorder.start()
    row = observation(sid)
    row["remaining"] = float("nan")
    try:
        with pytest.raises(StorageError, match="序列化"):
            recorder.submit("control", row)
        assert recorder.pending_count == 0
    finally:
        recorder.close({"session_id": sid}, START)
    assert store.query_controls(sid) == []


@pytest.mark.parametrize("suffix", [".txt", ".xlsx"])
def test_history_export_reads_snapshot_samples_and_controls_from_same_transaction(
    store, tmp_path, monkeypatch, suffix
):
    from contextlib import contextmanager

    sid = create(store)
    old_snapshot = {"session_id": sid, "marker": "old"}
    store.write_batch(
        samples=[sample(sid, 1)],
        controls=[observation(sid, 1, sequence=1)],
        snapshot=old_snapshot,
    )
    original_connection = store._connection
    committed = False

    @contextmanager
    def concurrent_connection():
        nonlocal committed
        with original_connection() as connection:

            def interleave_after_session_read(cursor, values):
                nonlocal committed
                row = sqlite3.Row(cursor, values)
                # 在会话行读取后立即完成另一连接的提交，稳定复现快照先于样本的竞争。
                if not committed and "snapshot" in row.keys() and "settings" in row.keys():
                    committed = True
                    store.write_batch(
                        samples=[sample(sid, 2)],
                        controls=[observation(sid, 2, sequence=2)],
                        snapshot={"session_id": sid, "marker": "new"},
                    )
                return row

            connection.row_factory = interleave_after_session_read
            yield connection

    monkeypatch.setattr(store, "_connection", concurrent_connection)
    destination = tmp_path / f"consistent{suffix}"
    store.export_session(sid, destination)
    assert committed
    tables = exported_tables(destination)
    assert dict(tables["快照信息"][1:])["marker"] == "old"
    assert len(tables["采样数据"]) == len(tables["控制器状态"]) == 2
    assert store.get_session(sid)["snapshot"]["marker"] == "new"
    assert len(store.query_samples(sid)) == len(store.query_controls(sid)) == 2
