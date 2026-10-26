import json
import sqlite3
import threading
import time

import pytest
from openpyxl import load_workbook

from cqu_monitor.storage import QueuedRecorder, SessionStore, StorageError


START = "2026-10-16T01:00:00+08:00"


def sample(sid, elapsed=0, cell=1, value=1.25, valid=True):
    return dict(
        session_id=sid,
        role="current" if cell is None else "voltage_left",
        cell=cell,
        round=1 if cell else None,
        phase=2,
        elapsed=elapsed,
        wall_time=START,
        value=value,
        valid=valid,
        status="ok",
        error=None,
        mean_value=value,
        sync_version=1,
    )


def snapshot(sid):
    return dict(
        session_id=sid,
        mode="simulation",
        elapsed=10,
        synchronized=True,
        sync_version=1,
        slot=1,
        round=1,
        phase=2,
        cell_pair=[1, 16],
        alarms=[],
        cells=[
            dict(cell=i, latest=i / 10, mean=i / 20, count=2, updated_at=10, age=0, status="ok")
            for i in range(1, 31)
        ],
        metrics=dict(
            total_voltage=46.5,
            avg_total_voltage=45.0,
            current=2.0,
            power_kw=0.093,
            avg_power_kw=0.09,
            complete=True,
            reason="complete",
        ),
    )


@pytest.fixture
def store(tmp_path):
    return SessionStore(tmp_path / "data" / "history.sqlite3")


def create(store):
    return store.create_session("simulation", {"simulation_speed": 10}, START)


def test_immediate_stop_drains_less_than_batch_and_final_snapshot(store):
    sid = create(store)
    recorder = QueuedRecorder(store, sid, batch_size=50)
    recorder.start()
    for i in range(13):
        recorder.submit("sample", sample(sid, i))
    recorder.close(snapshot(sid), START)
    assert len(store.query_samples(sid)) == 13
    assert recorder.pending_count == 0 and not recorder.alive
    assert store.get_session(sid)["snapshot"]["cells"][0]["count"] == 2
    assert store.get_session(sid)["status"] == "completed"


def test_periodic_flush_does_not_wait_for_fifty(store):
    sid = create(store)
    recorder = QueuedRecorder(store, sid, flush_interval=0.03)
    recorder.start()
    recorder.submit("sample", sample(sid))
    deadline = time.monotonic() + 2
    while recorder.pending_count and time.monotonic() < deadline:
        time.sleep(0.01)
    assert len(store.query_samples(sid)) == 1
    recorder.close(snapshot(sid), START)


def test_concurrent_producers_persist_every_accepted_sample(store):
    sid = create(store)
    recorder = QueuedRecorder(store, sid)
    recorder.start()
    errors = []

    def produce(base):
        try:
            for i in range(100):
                recorder.submit("sample", sample(sid, base + i))
        except Exception as exc:
            errors.append(exc)

    threads = [threading.Thread(target=produce, args=(i * 100,)) for i in range(4)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    recorder.close(snapshot(sid), START)
    assert not errors
    assert len(store.query_samples(sid)) == 400


def test_write_failure_preserves_failing_batch_and_queue(store, monkeypatch):
    sid = create(store)

    def disk_failure(**kwargs):
        raise OSError("injected disk full")

    monkeypatch.setattr(store, "write_batch", disk_failure)
    recorder = QueuedRecorder(store, sid, batch_size=50, flush_interval=30)
    recorder.start()
    for i in range(17):
        recorder.submit("sample", sample(sid, i))
    with pytest.raises(StorageError, match="disk full"):
        recorder.close(snapshot(sid), START)
    assert not recorder.alive and recorder.error
    records = [
        json.loads(line) for line in recorder.recovery_path.read_text(encoding="utf-8").splitlines()
    ]
    assert len([r for r in records if r.get("kind") == "sample"]) == 17
    assert any(r.get("kind") == "snapshot" for r in records)
    assert recorder.pending_count == 17
    with pytest.raises(StorageError):
        recorder.submit("sample", sample(sid))
    assert store.get_session(sid)["status"] != "completed"


def test_batch_is_atomic_and_sessions_isolated(store):
    first, second = create(store), create(store)
    store.write_batch(samples=[sample(first), sample(second, value=9)])
    with pytest.raises(StorageError):
        store.write_batch(samples=[sample(first, 1), sample("missing", 2)])
    assert [s["value"] for s in store.query_samples(first)] == [1.25]
    assert [s["value"] for s in store.query_samples(second)] == [9]


def test_filters_pagination_current_zero_and_unknown_ids(store):
    sid = create(store)
    store.write_batch(
        samples=[sample(sid, i, cell=None if i % 2 else 1, valid=i != 3) for i in range(8)]
    )
    assert [r["elapsed"] for r in store.query_samples(sid, start=2, end=6, limit=2, offset=1)] == [
        3,
        4,
    ]
    assert [r["elapsed"] for r in store.query_samples(sid, cell=0, valid_only=True)] == [1, 5, 7]
    assert len(store.query_samples(sid, start=START, end=START)) == 8
    assert store.query_samples("' OR 1=1 --") == []
    with pytest.raises(StorageError):
        store.get_session("' OR 1=1 --")
    with pytest.raises((ValueError, StorageError)):
        store.query_samples(sid, offset=-1)


def test_snapshot_and_history_excel_roundtrip_numeric_and_metadata(store, tmp_path):
    sid = create(store)
    summ = dict(
        session_id=sid,
        elapsed=3,
        wall_time=START,
        total_voltage=46.5,
        avg_total_voltage=45.0,
        current=2.0,
        power_kw=0.093,
        avg_power_kw=0.09,
        complete=True,
        reason="complete",
    )
    store.write_batch(
        samples=[sample(sid, 2), sample(sid, 3, None, 2)], summaries=[summ], snapshot=snapshot(sid)
    )
    history = tmp_path / "history.xlsx"
    store.export_session(sid, history, cell=0)
    wb = load_workbook(history, read_only=True, data_only=True)
    assert {"会话信息", "采样数据", "每秒汇总"}.issubset(wb.sheetnames)
    rows = list(wb["采样数据"].values)
    assert len(rows) == 2
    assert isinstance(rows[1][rows[0].index("value")], (int, float))
    assert rows[1][rows[0].index("value")] == 2
    assert "simulation" in str(list(wb["会话信息"].values))
    wb.close()
    current = tmp_path / "snapshot.xlsx"
    store.export_snapshot(snapshot(sid), current)
    wb = load_workbook(current, read_only=True, data_only=True)
    rows = list(wb["当前30室"].values)
    assert len(rows) == 31
    assert rows[1][:4] == (1, 0.1, 0.05, 2)
    assert "power_kw" in str(list(wb["底部指标"].values))
    wb.close()
    text_path = tmp_path / "history.txt"
    store.export_session(sid, text_path)
    assert text_path.read_bytes().startswith(b"\xef\xbb\xbf")
    assert "simulation" in text_path.read_text(encoding="utf-8-sig")


def test_export_rejects_missing_session_and_unsupported_suffix(store, tmp_path):
    sid = create(store)
    with pytest.raises(StorageError):
        store.export_session("missing", tmp_path / "missing.xlsx")
    assert not (tmp_path / "missing.xlsx").exists()
    with pytest.raises((ValueError, StorageError)):
        store.export_session(sid, tmp_path / "invalid.csv")


def test_stream_export_splits_before_excel_limit(store, tmp_path, monkeypatch):
    import cqu_monitor.storage as storage

    monkeypatch.setattr(storage, "EXCEL_MAX_ROWS", 4)
    sid = create(store)
    store.write_batch(samples=[sample(sid, i) for i in range(7)])
    dest = tmp_path / "split.xlsx"
    store.export_session(sid, dest, valid_only=False)
    wb = load_workbook(dest, read_only=True)
    sheets = [s for s in wb.sheetnames if s.startswith("采样数据")]
    assert len(sheets) == 3
    assert sum(len(list(wb[s].values)) - 1 for s in sheets) == 7
    wb.close()


def test_queue_full_is_visible_and_already_accepted_data_survives(store, monkeypatch):
    sid = create(store)
    original = store.write_batch
    entered, release = threading.Event(), threading.Event()

    def slow_write(**kwargs):
        entered.set()
        assert release.wait(2)
        original(**kwargs)

    monkeypatch.setattr(store, "write_batch", slow_write)
    recorder = QueuedRecorder(store, sid, maxsize=1, batch_size=1)
    recorder.start()
    recorder.submit("sample", sample(sid, 1))
    assert entered.wait(2)
    recorder.submit("sample", sample(sid, 2))
    with pytest.raises(StorageError, match="队列已满"):
        recorder.submit("sample", sample(sid, 3))
    release.set()
    recorder.close(snapshot(sid), START)
    assert [s["elapsed"] for s in store.query_samples(sid)] == [1, 2]


def test_close_timeout_cannot_later_mark_session_completed(store, monkeypatch):
    sid = create(store)
    original = store.write_batch
    entered, release = threading.Event(), threading.Event()

    def stalled_write(**kwargs):
        entered.set()
        assert release.wait(2)
        original(**kwargs)

    monkeypatch.setattr(store, "write_batch", stalled_write)
    recorder = QueuedRecorder(store, sid, batch_size=1)
    recorder.start()
    recorder.submit("sample", sample(sid))
    assert entered.wait(2)
    with pytest.raises(StorageError, match="超时"):
        recorder.close(snapshot(sid), START, timeout=0.01)
    release.set()
    recorder._thread.join(2)
    assert not recorder.alive
    assert store.get_session(sid)["status"] != "completed"
    assert recorder.recovery_path.exists()


def test_date_filters_compare_instants_across_timezone_offsets(store):
    sid = create(store)
    row = sample(sid)
    row["wall_time"] = "2026-10-15T17:00:00+00:00"
    store.write_batch(samples=[row])
    assert len(store.query_samples(sid, start=START, end=START)) == 1


def test_startup_marks_only_running_sessions_interrupted(store):
    running, completed = create(store), create(store)
    store.finish_session(completed, snapshot(completed), START)
    assert store.mark_interrupted_sessions() == 1
    assert store.get_session(running)["status"] == "interrupted"
    assert store.get_session(running)["ended_at"] is None
    assert store.get_session(completed)["status"] == "completed"


def test_close_between_flush_decision_and_empty_check_still_drains(store, monkeypatch):
    sid = create(store)
    recorder = QueuedRecorder(store, sid, batch_size=50, flush_interval=30)
    boundary, release = threading.Event(), threading.Event()
    original = recorder._closing.is_set

    def pause_after_observing_open():
        old_value = original()
        if (
            threading.current_thread() is recorder._thread
            and recorder._inflight
            and not old_value
            and not boundary.is_set()
        ):
            boundary.set()
            assert release.wait(2)
        return old_value

    monkeypatch.setattr(recorder._closing, "is_set", pause_after_observing_open)
    recorder.start()
    recorder.submit("sample", sample(sid))
    assert boundary.wait(2)
    errors = []

    def close():
        try:
            recorder.close(snapshot(sid), START)
        except Exception as exc:
            errors.append(exc)

    closer = threading.Thread(target=close)
    closer.start()
    assert recorder._closing.wait(2)
    release.set()
    closer.join(2)
    assert not closer.is_alive() and not errors
    assert len(store.query_samples(sid)) == 1
    assert recorder.pending_count == 0


def test_engine_invalid_clock_samples_persist_without_inventing_elapsed(store, tmp_path):
    from cqu_monitor.engine import Engine

    sid = create(store)
    engine = Engine({"mode": "hardware"}, sid)
    engine.start(0, START)
    invalid = engine.accept("current", 2, 0, float("nan"), START)
    assert invalid["elapsed"] is None and not invalid["valid"]
    recorder = QueuedRecorder(store, sid)
    recorder.start()
    recorder.submit("sample", invalid)
    recorder.submit("sample", sample(sid, 2))
    recorder.close(engine.snapshot(3), START)
    rows = store.query_samples(sid)
    assert [r["elapsed"] for r in rows] == [2, None]
    path = tmp_path / "invalid-clock.xlsx"
    store.export_session(sid, path, valid_only=False)
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["采样数据"].values)
    assert rows[-1][rows[0].index("elapsed")] is None
    wb.close()


def test_version_one_database_migrates_without_losing_samples(tmp_path):
    path = tmp_path / "old.sqlite3"
    with sqlite3.connect(path) as conn:
        conn.executescript("""
            CREATE TABLE sessions(id TEXT PRIMARY KEY,mode TEXT NOT NULL,
                started_at TEXT NOT NULL,ended_at TEXT,status TEXT NOT NULL,
                settings TEXT NOT NULL,snapshot TEXT);
            CREATE TABLE samples(id INTEGER PRIMARY KEY AUTOINCREMENT,
                session_id TEXT NOT NULL REFERENCES sessions(id),elapsed REAL NOT NULL,
                wall_time TEXT NOT NULL,cell INTEGER,role TEXT NOT NULL,
                valid INTEGER NOT NULL,payload TEXT NOT NULL);
            PRAGMA user_version=1;
        """)
        conn.execute(
            "INSERT INTO sessions VALUES(?,?,?,?,?,?,?)",
            ("legacy", "simulation", START, None, "running", "{}", None),
        )
        conn.execute(
            "INSERT INTO samples VALUES(?,?,?,?,?,?,?,?)",
            (1, "legacy", 1, START, 1, "voltage_left", 1, json.dumps(sample("legacy", 1))),
        )
    store = SessionStore(path)
    store.write_batch(samples=[sample("legacy", None, valid=False)])
    assert [row["elapsed"] for row in store.query_samples("legacy")] == [1, None]
    with sqlite3.connect(path) as conn:
        assert conn.execute("PRAGMA user_version").fetchone()[0] >= 2
