"""实验历史、流式导出和有界异步记录器；不依赖 Qt 或设备线程。"""

from __future__ import annotations

import csv
import json
import os
import queue
import sqlite3
import tempfile
import threading
import time
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from numbers import Real
from pathlib import Path


SCHEMA_VERSION = 3
EXCEL_MAX_ROWS = 1_048_576
EXPORT_BATCH_SIZE = 1000
SAMPLE_FIELDS = (
    "session_id",
    "role",
    "cell",
    "round",
    "phase",
    "elapsed",
    "wall_time",
    "value",
    "valid",
    "status",
    "error",
    "mean_value",
    "sync_version",
    "synchronization_source",
    "sync_evidence",
)
CONTROL_FIELDS = (
    "session_id",
    "wall_time",
    "elapsed",
    "group",
    "sequence",
    "channel",
    "remaining",
    "width",
    "outputs",
    "cycle",
    "interval_remaining",
    "cycle_remaining",
    "status",
    "error",
    "started",
    "finished",
    "raw_parameters",
    "raw_registers",
)
SUMMARY_FIELDS = (
    "session_id",
    "elapsed",
    "wall_time",
    "total_voltage",
    "avg_total_voltage",
    "current",
    "power_kw",
    "avg_power_kw",
    "complete",
    "reason",
    "measured_count",
    "available_count",
    "coverage",
    "quality",
)
CALCULATION_NOTES = (
    "缺测为空，不补0；每室保存全部有效样本，显示最新值，均值取本实验全部有效样本。"
    "总压从首个已测小室开始，取所有已测小室最新有效电压之和；新值替换旧值。"
    "未齐全时为已测合计，含过期/故障设备旧值时保留合计并标注质量。功率kW=当前合计V×有效电流A/1000；"
    "总压均值及平均功率按引擎固定实验秒网格统计。"
)


class StorageError(RuntimeError):
    """需要向操作者明确报告的持久化或导出错误。"""


def _json(value):
    # 拒绝 NaN/Infinity，避免数据库与导出文件悄悄出现不可移植的缺测值。
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def _excel_value(value):
    if isinstance(value, (dict, list, tuple)):
        return _json(value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _unit(sample):
    return "A" if sample.get("role") == "current" else "V"


class SessionStore:
    """连接只在当前操作中使用；后台写入与历史查询不共享 SQLite 连接。"""

    def __init__(self, db_path: Path):
        self.db_path = Path(db_path).resolve()
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        with self._connection() as connection:
            version = connection.execute("PRAGMA user_version").fetchone()[0]
            if version not in (0, 1, 2, SCHEMA_VERSION):
                raise StorageError(f"数据库版本 {version} 不受支持，当前版本为 {SCHEMA_VERSION}")
            connection.execute("PRAGMA journal_mode=WAL")
            if version == 1:
                # 首版把elapsed设为非空，但引擎必须保留非法时钟的原始失败记录。
                # 重建表只放宽该约束，原编号/载荷全部保留，绝不把缺时钟补成0。
                connection.execute("BEGIN IMMEDIATE")
                connection.execute(
                    "CREATE TABLE samples_v2 (id INTEGER PRIMARY KEY AUTOINCREMENT,"
                    "session_id TEXT NOT NULL REFERENCES sessions(id),elapsed REAL,"
                    "wall_time TEXT NOT NULL,cell INTEGER,role TEXT NOT NULL,"
                    "valid INTEGER NOT NULL,payload TEXT NOT NULL)"
                )
                connection.execute("INSERT INTO samples_v2 SELECT * FROM samples")
                connection.execute("DROP TABLE samples")
                connection.execute("ALTER TABLE samples_v2 RENAME TO samples")
                connection.execute("PRAGMA user_version=2")
                connection.commit()
            connection.executescript(
                """
                BEGIN IMMEDIATE;
                CREATE TABLE IF NOT EXISTS sessions (
                    id TEXT PRIMARY KEY, mode TEXT NOT NULL, started_at TEXT NOT NULL,
                    ended_at TEXT, status TEXT NOT NULL, settings TEXT NOT NULL,
                    snapshot TEXT
                );
                CREATE TABLE IF NOT EXISTS samples (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    session_id TEXT NOT NULL REFERENCES sessions(id),
                    elapsed REAL, wall_time TEXT NOT NULL,
                    cell INTEGER, role TEXT NOT NULL, valid INTEGER NOT NULL,
                    payload TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS summaries (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    session_id TEXT NOT NULL REFERENCES sessions(id),
                    elapsed REAL NOT NULL, wall_time TEXT NOT NULL, payload TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS events (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    session_id TEXT NOT NULL REFERENCES sessions(id),
                    wall_time TEXT NOT NULL, payload TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS control_observations (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    session_id TEXT NOT NULL REFERENCES sessions(id),
                    elapsed REAL, wall_time TEXT NOT NULL,
                    group_name TEXT NOT NULL CHECK(group_name IN ('left','right')),
                    sequence INTEGER NOT NULL, payload TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS controls_session_elapsed
                    ON control_observations(session_id, elapsed, id);
                CREATE INDEX IF NOT EXISTS controls_session_wall
                    ON control_observations(session_id, wall_time);
                CREATE INDEX IF NOT EXISTS samples_session_elapsed
                    ON samples(session_id, elapsed, id);
                CREATE INDEX IF NOT EXISTS samples_session_wall
                    ON samples(session_id, wall_time);
                CREATE INDEX IF NOT EXISTS summaries_session_elapsed
                    ON summaries(session_id, elapsed, id);
                CREATE INDEX IF NOT EXISTS summaries_session_wall
                    ON summaries(session_id, wall_time);
                CREATE INDEX IF NOT EXISTS events_session ON events(session_id, id);
                PRAGMA user_version=3;
                COMMIT;
                """
            )

    @contextmanager
    def _connection(self):
        """为当前操作建立独立SQLite连接，并在作用域结束时可靠关闭。"""
        connection = None
        try:
            connection = sqlite3.connect(self.db_path, timeout=5)
            connection.row_factory = sqlite3.Row
            connection.execute("PRAGMA busy_timeout=5000")
            connection.execute("PRAGMA foreign_keys=ON")
            with connection:
                yield connection
        except (sqlite3.Error, OSError, TypeError, ValueError, KeyError) as exc:
            raise StorageError(f"数据保存/读取失败：{exc}") from exc
        finally:
            if connection is not None:
                connection.close()

    def create_session(self, mode, settings, started_at, session_id=None):
        """建立带完整参数的新会话，样本和控制证据均通过此编号关联。"""
        session_id = session_id or uuid.uuid4().hex
        with self._connection() as connection:
            connection.execute(
                "INSERT INTO sessions(id,mode,started_at,status,settings) VALUES(?,?,?,?,?)",
                (session_id, mode, started_at, "running", _json(settings)),
            )
        return session_id

    def write_batch(self, samples=(), summaries=(), events=(), snapshot=None, controls=()):
        # 四类记录及最新快照处于同一个事务：任何一行失败，整批回滚，交给队列恢复。
        """在一个事务内写入本批样本、汇总、事件和控制状态；失败整批回滚。"""
        with self._connection() as connection:
            connection.executemany(
                "INSERT INTO samples(session_id,elapsed,wall_time,cell,role,valid,payload) "
                "VALUES(?,?,?,?,?,?,?)",
                (
                    (
                        s["session_id"],
                        s["elapsed"],
                        s["wall_time"],
                        s.get("cell"),
                        s["role"],
                        bool(s["valid"]),
                        _json(s),
                    )
                    for s in samples
                ),
            )
            connection.executemany(
                "INSERT INTO summaries(session_id,elapsed,wall_time,payload) VALUES(?,?,?,?)",
                ((s["session_id"], s["elapsed"], s["wall_time"], _json(s)) for s in summaries),
            )
            connection.executemany(
                "INSERT INTO events(session_id,wall_time,payload) VALUES(?,?,?)",
                ((s["session_id"], s["wall_time"], _json(s)) for s in events),
            )
            connection.executemany(
                "INSERT INTO control_observations"
                "(session_id,elapsed,wall_time,group_name,sequence,payload) VALUES(?,?,?,?,?,?)",
                (
                    (
                        s["session_id"],
                        s["elapsed"],
                        s["wall_time"],
                        s["group"],
                        s["sequence"],
                        _json(s),
                    )
                    for s in controls
                ),
            )
            if snapshot is not None:
                cursor = connection.execute(
                    "UPDATE sessions SET snapshot=? WHERE id=?",
                    (_json(snapshot), snapshot["session_id"]),
                )
                if cursor.rowcount != 1:
                    raise StorageError(f"实验不存在：{snapshot['session_id']}")

    def finish_session(self, session_id, snapshot, ended_at, status="completed"):
        """保存最终快照、结束时间及最终状态，不删除既有采样明细。"""
        if snapshot is not None and snapshot.get("session_id") != session_id:
            raise StorageError("最终快照与实验编号不一致")
        with self._connection() as connection:
            cursor = connection.execute(
                "UPDATE sessions SET snapshot=COALESCE(?,snapshot), ended_at=?,status=? WHERE id=?",
                (_json(snapshot) if snapshot is not None else None, ended_at, status, session_id),
            )
            if cursor.rowcount != 1:
                raise StorageError(f"实验不存在：{session_id}")

    def record_session_errors(self, session_id, events):
        """在同一事务中补记保存阶段晚到的异常并更正会话状态，不改测量数据。"""
        events = list(events)
        if not events or any(event.get("session_id") != session_id for event in events):
            raise StorageError("补记异常事件必须属于同一实验")
        with self._connection() as connection:
            cursor = connection.execute(
                "UPDATE sessions SET status='error' WHERE id=? AND ended_at IS NOT NULL",
                (session_id,),
            )
            if cursor.rowcount != 1:
                raise StorageError("实验尚未完成写库，不能提前补写最终状态")
            connection.executemany(
                "INSERT INTO events(session_id,wall_time,payload) VALUES(?,?,?)",
                ((session_id, event["wall_time"], _json(event)) for event in events),
            )

    def mark_interrupted_sessions(self):
        """仅由应用启动流程调用；打开历史数据库本身不会终止其他在采实验。"""
        with self._connection() as connection:
            cursor = connection.execute(
                "UPDATE sessions SET status='interrupted' WHERE status='running'"
            )
            return cursor.rowcount

    def list_sessions(self):
        """按实际开始时刻倒序读取会话；相同时刻按较晚建档优先，不修改记录。"""
        with self._connection() as connection:
            return [
                dict(row)
                for row in connection.execute(
                    "SELECT id,mode,started_at,ended_at,status FROM sessions "
                    # ISO中的本地钟面可能更晚，但时区换算后的实际时刻更早。
                    "ORDER BY julianday(started_at) DESC,rowid DESC"
                )
            ]

    @staticmethod
    def _read_session(connection, session_id):
        row = connection.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
        if row is None:
            raise StorageError(f"实验不存在：{session_id}")
        result = dict(row)
        result["settings"] = json.loads(result["settings"])
        result["snapshot"] = json.loads(result["snapshot"]) if result["snapshot"] else None
        return result

    def get_session(self, session_id):
        """读取单个会话及解码快照；不存在时明确报错而非构造空实验。"""
        with self._connection() as connection:
            return self._read_session(connection, session_id)

    def clear_history(self):
        """一个写事务清除父子记录；运行中拒绝，任一步失败即整批回滚。"""
        with self._connection() as connection:
            connection.execute("BEGIN IMMEDIATE")
            if connection.execute(
                "SELECT 1 FROM sessions WHERE status='running' LIMIT 1"
            ).fetchone():
                raise StorageError("存在运行中的实验，请停止采集后再清空历史记录")
            count = connection.execute("SELECT COUNT(*) FROM sessions").fetchone()[0]
            for table in ("samples", "summaries", "events", "control_observations", "sessions"):
                connection.execute(f"DELETE FROM {table}")
        return count

    @staticmethod
    def _filters(session_id, start, end, cell=None, valid_only=False):
        clauses, values = ["session_id=?"], [session_id]
        for boundary, operator in ((start, ">="), (end, "<=")):
            if boundary is not None:
                if isinstance(boundary, Real):
                    clauses.append(f"elapsed{operator}?")
                else:
                    # ISO字符串可能带不同时区；比较同一时刻，而不是比较日期文字。
                    clauses.append(f"julianday(wall_time){operator}julianday(?)")
                values.append(boundary)
        if cell is not None:
            if not isinstance(cell, int) or not 0 <= cell <= 30:
                raise StorageError("小室筛选必须为0（电流）或1至30")
            if cell == 0:
                clauses.append("role=?")
                values.append("current")
            else:
                clauses.append("cell=?")
                values.append(cell)
        if valid_only:
            clauses.append("valid=1")
        return " AND ".join(clauses), values

    def _query(self, table, session_id, start, end, limit, offset, cell=None, valid_only=False):
        if not isinstance(limit, int) or not isinstance(offset, int) or limit < 0 or offset < 0:
            raise StorageError("分页数量与偏移必须为非负整数")
        where, values = self._filters(session_id, start, end, cell, valid_only)
        with self._connection() as connection:
            rows = connection.execute(
                f"SELECT payload FROM {table} WHERE {where} "
                "ORDER BY elapsed IS NULL,elapsed,id LIMIT ? OFFSET ?",
                (*values, limit, offset),
            )
            return [json.loads(row[0]) for row in rows]

    def query_samples(
        self, session_id, start=None, end=None, cell=None, valid_only=False, limit=1000, offset=0
    ):
        return self._query("samples", session_id, start, end, limit, offset, cell, valid_only)

    def query_summaries(self, session_id, start=None, end=None, limit=1000, offset=0):
        return self._query("summaries", session_id, start, end, limit, offset)

    @classmethod
    def _control_filters(cls, session_id, start, end, group=None):
        where, values = cls._filters(session_id, start, end)
        if group is not None:
            if group not in ("left", "right"):
                raise StorageError("控制器组筛选必须为left或right")
            where += " AND group_name=?"
            values.append(group)
        return where, values

    def query_controls(self, session_id, start=None, end=None, group=None, limit=1000, offset=0):
        if not isinstance(limit, int) or not isinstance(offset, int) or limit < 0 or offset < 0:
            raise StorageError("分页数量与偏移必须为非负整数")
        where, values = self._control_filters(session_id, start, end, group)
        with self._connection() as connection:
            rows = connection.execute(
                f"SELECT payload FROM control_observations WHERE {where} "
                "ORDER BY elapsed IS NULL,elapsed,id LIMIT ? OFFSET ?",
                (*values, limit, offset),
            )
            return [json.loads(row[0]) for row in rows]

    @staticmethod
    def _stream(connection, table, where, values, order="elapsed IS NULL,elapsed,id"):
        cursor = connection.execute(
            f"SELECT payload FROM {table} WHERE {where} ORDER BY {order}", values
        )
        while batch := cursor.fetchmany(EXPORT_BATCH_SIZE):
            for row in batch:
                yield json.loads(row[0])

    @staticmethod
    def _metadata(session, filters=None):
        rows = [
            (key, session.get(key))
            for key in ("id", "mode", "started_at", "ended_at", "status", "settings")
        ]
        rows.extend(
            [
                ("计算口径", CALCULATION_NOTES),
                ("电压单位", "V"),
                ("电流单位", "A"),
                ("功率单位", "kW"),
                ("筛选条件", filters or {}),
                ("schema_version", SCHEMA_VERSION),
            ]
        )
        return rows

    def export_session(
        self, session_id, destination, start=None, end=None, cell=None, valid_only=True
    ):
        # 会话快照和所有数据流共用同一读事务；先读取会话固定WAL视图，
        # 随后的后台提交不能把新版样本/控制证据混入旧版快照。
        """在同一个WAL读视图流式导出各类明细，避免边采集边导出时跨版本混合。"""
        with self._connection() as connection:
            connection.execute("BEGIN")
            session = self._read_session(connection, session_id)
            filters = dict(start=start, end=end, cell=cell, valid_only=valid_only)
            sample_where, sample_values = self._filters(session_id, start, end, cell, valid_only)
            summary_where, summary_values = self._filters(session_id, start, end)
            # 电压按小室归属组导出；电流独立采集，保留同期左右控制器上下文。
            # valid_only仅筛选仪表样本，控制器错误状态也是必须保留的审计证据。
            control_group = None if cell in (None, 0) else ("left" if cell <= 15 else "right")
            control_where, control_values = self._control_filters(
                session_id, start, end, control_group
            )
            with _ExportFile(destination) as writer:
                writer.table("会话信息", ("项目", "值"), self._metadata(session, filters))
                writer.table(
                    "采样数据",
                    (*SAMPLE_FIELDS, "unit"),
                    (
                        [*(s.get(k) for k in SAMPLE_FIELDS), _unit(s)]
                        for s in self._stream(connection, "samples", sample_where, sample_values)
                    ),
                )
                writer.table(
                    "每秒汇总",
                    SUMMARY_FIELDS,
                    (
                        [s.get(k) for k in SUMMARY_FIELDS]
                        for s in self._stream(
                            connection, "summaries", summary_where, summary_values
                        )
                    ),
                )
                writer.table(
                    "控制器状态",
                    (*CONTROL_FIELDS, "payload"),
                    (
                        [*(s.get(k) for k in CONTROL_FIELDS), s]
                        for s in self._stream(
                            connection, "control_observations", control_where, control_values
                        )
                    ),
                )
                # 事件没有 elapsed，按实际墙钟筛选；数值 elapsed 过滤不虚构事件时间。
                event_start = start if isinstance(start, str) else None
                event_end = end if isinstance(end, str) else None
                event_where, event_values = self._filters(session_id, event_start, event_end)
                writer.table(
                    "事件记录",
                    ("session_id", "wall_time", "level", "message"),
                    (
                        [s.get(k) for k in ("session_id", "wall_time", "level", "message")]
                        for s in self._stream(
                            connection, "events", event_where, event_values, order="id"
                        )
                    ),
                )
                if session["snapshot"]:
                    self._snapshot_tables(writer, session["snapshot"])
        return Path(destination)

    @staticmethod
    def _snapshot_tables(writer, snapshot):
        fields = ("cell", "latest", "mean", "count", "unit", "updated_at", "age", "status")
        cells = snapshot.get("cells", [])
        writer.table(
            "当前30室", fields, ([c.get(k) if k != "unit" else "V" for k in fields] for c in cells)
        )
        metrics = snapshot.get("metrics", {})
        units = {
            "total_voltage": "V",
            "avg_total_voltage": "V",
            "current": "A",
            "power_kw": "kW",
            "avg_power_kw": "kW",
        }
        writer.table(
            "底部指标",
            ("metric", "value", "unit"),
            ((key, value, units.get(key, "")) for key, value in metrics.items()),
        )
        writer.table(
            "快照信息",
            ("项目", "值"),
            ((key, value) for key, value in snapshot.items() if key not in ("cells", "metrics")),
        )

    def export_snapshot(self, snapshot, destination):
        """导出固定30室及底部指标；缺测为留空，输出携带原会话和计算口径。"""
        if not isinstance(snapshot, dict) or not snapshot.get("session_id"):
            raise StorageError("没有可导出的实验快照")
        if len(snapshot.get("cells", [])) != 30:
            raise StorageError("快照必须包含30个小室（缺测值使用空值）")
        session = self.get_session(snapshot["session_id"])
        with _ExportFile(destination) as writer:
            writer.table("会话信息", ("项目", "值"), self._metadata(session))
            self._snapshot_tables(writer, snapshot)
        return Path(destination)


class _ExportFile:
    """先写同目录临时文件，成功关闭后替换目标，失败不留下伪装完整的导出。"""

    def __init__(self, destination):
        self.destination = Path(destination)
        self.suffix = self.destination.suffix.lower()
        if self.suffix not in (".xlsx", ".txt"):
            raise StorageError("导出格式仅支持 .xlsx 或 .txt")
        self.workbook = None
        self.handle = None
        self.temp_path = None

    def __enter__(self):
        try:
            self.destination.parent.mkdir(parents=True, exist_ok=True)
            fd, name = tempfile.mkstemp(
                suffix=self.suffix, prefix=".export-", dir=self.destination.parent
            )
            os.close(fd)
            self.temp_path = Path(name)
            if self.suffix == ".xlsx":
                from openpyxl import Workbook

                self.workbook = Workbook(write_only=True)
            else:
                self.handle = self.temp_path.open("w", encoding="utf-8-sig", newline="")
                self.text_writer = csv.writer(self.handle, delimiter="\t")
            return self
        except Exception as exc:
            self._cleanup()
            raise StorageError(f"无法创建导出文件：{exc}") from exc

    def table(self, name, headers, rows):
        """流式写一张表；按后缀输出UTF-8文本或真正xlsx，不把文本伪装成Excel。"""
        if self.workbook is not None:
            sheet_number, count = 1, 1
            sheet = self.workbook.create_sheet(name)
            sheet.append(headers)
            for row in rows:
                if count >= EXCEL_MAX_ROWS:
                    sheet_number += 1
                    sheet = self.workbook.create_sheet(f"{name}_{sheet_number}")
                    sheet.append(headers)
                    count = 1
                sheet.append([_excel_value(value) for value in row])
                count += 1
        else:
            self.text_writer.writerow([f"[{name}]"])
            self.text_writer.writerow(headers)
            for row in rows:
                self.text_writer.writerow(
                    [_json(v) if isinstance(v, (dict, list, tuple)) else v for v in row]
                )
            self.text_writer.writerow([])

    def _cleanup(self):
        if self.handle:
            self.handle.close()
        if self.workbook:
            self.workbook.close()
        if self.temp_path and self.temp_path.exists():
            self.temp_path.unlink()

    def __exit__(self, exc_type, exc, traceback):
        try:
            if exc_type is None:
                if self.workbook:
                    self.workbook.save(self.temp_path)
                    self.workbook.close()
                if self.handle:
                    self.handle.flush()
                    os.fsync(self.handle.fileno())
                    self.handle.close()
                os.replace(self.temp_path, self.destination)
        except Exception as failure:
            raise StorageError(f"导出失败：{failure}") from failure
        finally:
            self._cleanup()
        return False


class QueuedRecorder:
    """只有后台线程写库；只有提交成功后才减去待写数量。"""

    def __init__(self, store, session_id, maxsize=10000, batch_size=50, flush_interval=0.5):
        if maxsize < 1 or batch_size < 1 or flush_interval <= 0:
            raise ValueError("队列容量、批量及刷新间隔必须大于0")
        store.get_session(session_id)
        self.store, self.session_id = store, session_id
        self.batch_size, self.flush_interval = batch_size, flush_interval
        self._queue = queue.Queue(maxsize=maxsize)
        self._lock = threading.RLock()
        self._recovery_lock = threading.Lock()
        self._closing = threading.Event()
        self._thread = None
        self._error = None
        self._pending = 0
        self._inflight = []
        self._failed_items = []
        self._final_snapshot = None
        self._ended_at = None
        self._status = "completed"
        self._closed = False
        self._recovery_path = None
        self._recovery_written = False

    @property
    def error(self):
        with self._lock:
            return self._error

    @property
    def pending_count(self):
        with self._lock:
            return self._pending

    @property
    def alive(self):
        return self._thread is not None and self._thread.is_alive()

    @property
    def recovery_path(self):
        return self._recovery_path

    def start(self):
        """启动唯一写库线程；主线程只负责提交，设备进程不直接接触数据库。"""
        with self._lock:
            if self._thread is not None or self._closed:
                raise StorageError("记录器已经启动或关闭")
            self._thread = threading.Thread(target=self._run, name="session-recorder", daemon=True)
            self._thread.start()

    def submit(self, kind, payload):
        """短锁内提交有界任务；关闭、故障或队列满时拒绝并由调用方保留载荷。"""
        if kind not in ("sample", "summary", "event", "snapshot", "control"):
            raise StorageError(f"未知记录类型：{kind}")
        if payload.get("session_id") != self.session_id:
            raise StorageError("记录与当前实验编号不一致")
        # 入队即冻结数据；调用方之后更新界面快照不能改变尚未写库的历史。
        try:
            frozen = json.loads(_json(payload))
        except (TypeError, ValueError) as exc:
            raise StorageError(f"记录不能序列化：{exc}") from exc
        with self._lock:
            if self._error:
                raise StorageError(self._error)
            if self._thread is None or self._closing.is_set() or not self.alive:
                raise StorageError("记录器尚未启动或已经停止")
            try:
                self._queue.put_nowait((kind, frozen))
            except queue.Full as exc:
                # 满队列明确拒绝本条，由控制器停止生产；不冒充接受，不丢弃既有数据。
                raise StorageError("保存队列已满，请停止采集并排空已有数据") from exc
            self._pending += 1

    def _flush(self):
        """提交内存批次，只有事务成功才减少pending计数；失败保留在途原包。"""
        items = self._inflight
        if not items:
            return
        snapshots = [payload for kind, payload in items if kind == "snapshot"]
        controls = [payload for kind, payload in items if kind == "control"]
        # 无控制状态时保留旧记录器接口兼容性；有状态才向存储传入新增参数。
        control_kwargs = {"controls": controls} if controls else {}
        self.store.write_batch(
            samples=[p for k, p in items if k == "sample"],
            summaries=[p for k, p in items if k == "summary"],
            events=[p for k, p in items if k == "event"],
            snapshot=snapshots[-1] if snapshots else None,
            **control_kwargs,
        )
        with self._lock:
            self._pending -= len(items)
            self._inflight = []

    def _record_failure(self, failure):
        """冻结首次保存错误并收齐未提交任务，再在共享状态锁外写恢复材料。"""
        with self._lock:
            if not self._error:
                self._error = f"实验保存失败：{failure}"
            self._closing.set()
            # 未提交批次保留原顺序，先拒绝生产者，再收走剩余队列，避免恢复时遗漏尾部。
            if not self._failed_items:
                self._failed_items = list(self._inflight)
                while True:
                    try:
                        self._failed_items.append(self._queue.get_nowait())
                    except queue.Empty:
                        break
        # 恢复文件I/O不能持有状态锁；界面读取error/submit不等待坏盘。
        self._save_recovery()

    def _save_recovery(self, force=False):
        # 单独序列化恢复写盘；短状态锁只冻结本轮载荷，不包住open/write/fsync。
        """单独串行写恢复快照：临时文件写完并fsync后替换，失败保护上一完整副本。"""
        with self._recovery_lock:
            with self._lock:
                if self._recovery_written and not force:
                    return
                directory = self.store.db_path.parent / "recovery"
                if self._recovery_path is None:
                    self._recovery_path = directory / f"recovery-{uuid.uuid4().hex}.jsonl"
                metadata = dict(
                    kind="recovery_metadata",
                    session_id=self.session_id,
                    error=self._error,
                    ended_at=self._ended_at,
                    requested_status=self._status,
                    created_at=datetime.now(timezone.utc).isoformat(),
                )
                records = list(self._failed_items)
                snapshot = self._final_snapshot
            temporary = None
            try:
                directory.mkdir(parents=True, exist_ok=True)
                descriptor, name = tempfile.mkstemp(
                    prefix="recovery-", suffix=".tmp", dir=directory
                )
                temporary = Path(name)
                with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
                    handle.write(_json(metadata) + "\n")
                    for kind, payload in records:
                        handle.write(_json(dict(kind=kind, payload=payload)) + "\n")
                    if snapshot is not None:
                        handle.write(_json(dict(kind="snapshot", payload=snapshot)) + "\n")
                    handle.flush()
                    os.fsync(handle.fileno())
                os.replace(temporary, self._recovery_path)
                with self._lock:
                    self._recovery_written = True
            except Exception as exc:
                with self._lock:
                    self._error = (
                        f"{self._error}；恢复文件写入也失败：{exc}；未提交数据仍保留在记录器内存"
                    )
            finally:
                if temporary is not None:
                    try:
                        temporary.unlink(missing_ok=True)
                    except OSError:
                        pass

    def _run(self):
        """后台按批次/时间提交；关闭信号到达后排完内存批次和队列，最后结束会话。"""
        last_flush = time.monotonic()
        try:
            while True:
                if self.error:
                    return
                timeout = max(
                    0.001, min(0.05, self.flush_interval - (time.monotonic() - last_flush))
                )
                got_item = False
                with self._lock:
                    try:
                        item = self._queue.get_nowait()
                        self._inflight.append(item)
                        got_item = True
                    except queue.Empty:
                        pass
                if not got_item:
                    self._closing.wait(timeout)
                due = time.monotonic() - last_flush >= self.flush_interval
                if len(self._inflight) >= self.batch_size or due or self._closing.is_set():
                    self._flush()
                    last_flush = time.monotonic()
                # close超时可能发生在一次已开始的写库中；即便该事务后来成功，
                # 也不能把一次已经向用户报告失败的退出改写成正常结束。
                if self.error:
                    return
                if self._closing.is_set() and self._queue.empty():
                    # 停止信号可能恰好到达上面的批量判断之后：队列空不等于
                    # 已取出的内存批次已提交，结束实验前必须再排空一次。
                    self._flush()
                    if self.error:
                        return
                    self.store.finish_session(
                        self.session_id, self._final_snapshot, self._ended_at, self._status
                    )
                    return
        except Exception as exc:
            self._record_failure(exc)

    def close(self, final_snapshot, ended_at, status="completed", timeout=15):
        """仅在所有生产者退出后调用；排空写库，失败时保存恢复证据并抛StorageError。"""
        if final_snapshot is not None and final_snapshot.get("session_id") != self.session_id:
            raise StorageError("最终快照与当前实验编号不一致")
        with self._lock:
            if self._closed and not self._error:
                return
            if self._thread is None:
                raise StorageError("记录器尚未启动")
            try:
                self._final_snapshot = json.loads(_json(final_snapshot))
            except (TypeError, ValueError) as exc:
                raise StorageError(f"最终快照不能序列化：{exc}") from exc
            self._ended_at, self._status = ended_at, status
            self._closing.set()
        self._thread.join(max(0, timeout))
        if self.alive:
            self._record_failure("等待写库线程退出超时，正在进行的数据库事务结果尚未确定")
        with self._lock:
            self._closed = True
            failed = bool(self._error)
        if failed:
            # 强制重写可覆盖另一恢复写入者的旧快照，但不长时间持有状态锁。
            self._save_recovery(force=True)
            raise StorageError(f"{self.error}；恢复文件：{self._recovery_path}")
